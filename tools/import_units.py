#!/usr/bin/env python3
"""
Converts content/units.xlsx (one sheet per faction: Lp. | Name | Points) and content/items.xlsx
(upgrades/items: "Neutral Upgrades" for every faction + one sheet per faction) into
backend/src/main/resources/content/eldfall-armies.json used by the warband builder.

Usage (from the repo root):  python3 tools/import_units.py
Requires: pip install openpyxl

Characters with the same name in several sheets (e.g. Adventurers' Guild members that also belong to
Soga/Helian/Thenion/Sand) are one character with one code; their points must be identical.
The same holds for items listed in several faction sheets (e.g. Seasoned Combatant).
Costs like "2 pkt / 1 pkt*" become points=2, reducedPoints=1 with the condition in REDUCED_NOTES.
"""
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "content" / "units.xlsx"
ITEMS_SRC = ROOT / "content" / "items.xlsx"
OUT = ROOT / "backend" / "src" / "main" / "resources" / "content" / "eldfall-armies.json"

# Sheet name -> faction code (same codes as the scheme content).
SHEETS = {
    "Coalition of Thenion": "COALITION_OF_THENION",
    "Empire of Soga": "EMPIRE_OF_SOGA",
    "Helian League": "HELIAN_LEAGUE",
    "Sand Kingdoms": "SAND_KINGDOMS",
    "Oni Clans": "ONI_CLANS",
    "Goblin Wartribes": "GOBLIN_WARTRIBES",
    "Adventures Guild": "ADVENTURERS_GUILD",
    "Undead": "UNDEAD",
}

# Who may be played as the main faction, and which extra lists it may recruit from.
FACTIONS = [
    {"code": "COALITION_OF_THENION", "name": "Coalition of Thenion", "playable": True,
     "alwaysAvailable": ["ADVENTURERS_GUILD"], "optionalAllies": []},
    {"code": "EMPIRE_OF_SOGA", "name": "Empire of Soga", "playable": True,
     "alwaysAvailable": ["ADVENTURERS_GUILD"], "optionalAllies": []},
    {"code": "HELIAN_LEAGUE", "name": "Helian League", "playable": True,
     "alwaysAvailable": ["ADVENTURERS_GUILD"], "optionalAllies": []},
    {"code": "SAND_KINGDOMS", "name": "Sand Kingdoms", "playable": True,
     "alwaysAvailable": ["ADVENTURERS_GUILD"], "optionalAllies": []},
    {"code": "ONI_CLANS", "name": "Oni Clans", "playable": True,
     "alwaysAvailable": [], "optionalAllies": ["UNDEAD"]},
    {"code": "GOBLIN_WARTRIBES", "name": "Goblin Wartribes", "playable": True,
     "alwaysAvailable": [], "optionalAllies": []},
    {"code": "ADVENTURERS_GUILD", "name": "Adventurers Guild", "playable": True,
     "alwaysAvailable": [], "optionalAllies": []},
    {"code": "UNDEAD", "name": "Undead", "playable": False,
     "alwaysAvailable": [], "optionalAllies": []},
]


ITEM_SHEETS = {
    "Neutral Upgrades": "NEUTRAL",
    "Coalition of Thenion": "COALITION_OF_THENION",
    "Empire of Soga": "EMPIRE_OF_SOGA",
    "Helian League": "HELIAN_LEAGUE",
    "Sand Kingdoms": "SAND_KINGDOMS",
}

# Conditions for the reduced cost (from the footnotes in items.xlsx).
REDUCED_NOTES = {
    "KASSEN_BUKI_YARI_TSUKAI": {"pl": "jeśli broń podstawowa to Spear", "en": "if the primary weapon is a Spear"},
    "KASSEN_BUKI_DAIKYUU_TSUKAI": {"pl": "jeśli postać ma Short Bow", "en": "if the character has a Short Bow"},
}


def read_items() -> tuple[list[dict], dict[str, list[str]]]:
    wb = openpyxl.load_workbook(ITEMS_SRC, data_only=True)
    items: dict[str, dict] = {}
    lists: dict[str, list[str]] = {}
    for sheet, key in ITEM_SHEETS.items():
        if sheet not in wb.sheetnames:
            raise SystemExit(f"missing item sheet: {sheet}")
        codes = []
        for row in wb[sheet].iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None or not row[1] or row[2] is None:
                continue  # footnotes have text only in the first column
            name = str(row[1]).strip()
            costs = [int(x) for x in re.findall(r"\d+", str(row[2]))]
            code = code_of(name)
            item = {"code": code, "name": name, "points": costs[0]}
            if len(costs) > 1:
                if code not in REDUCED_NOTES:
                    raise SystemExit(f"{name}: two costs but no REDUCED_NOTES entry")
                item["reducedPoints"] = costs[1]
                item["reducedNote"] = REDUCED_NOTES[code]
            if code in items and items[code] != item:
                raise SystemExit(f"{name}: different costs in two sheets")
            items.setdefault(code, item)
            if code not in codes:
                codes.append(code)
        lists[key] = codes
    return sorted(items.values(), key=lambda i: i["name"]), lists


def code_of(name: str) -> str:
    ascii_name = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]+", "_", ascii_name.upper()).strip("_")


def main() -> int:
    wb = openpyxl.load_workbook(SRC, data_only=True)
    units: dict[str, dict] = {}
    lists: dict[str, list[str]] = {}
    for sheet, faction in SHEETS.items():
        if sheet not in wb.sheetnames:
            print(f"missing sheet: {sheet}", file=sys.stderr)
            return 1
        codes = []
        for row in wb[sheet].iter_rows(min_row=2, values_only=True):
            if not row or not row[1]:
                continue
            name = str(row[1]).strip()
            points = int(re.sub(r"[^0-9]", "", str(row[2])))
            code = code_of(name)
            if code in units and units[code]["points"] != points:
                print(f"{name}: different points ({units[code]['points']} vs {points})", file=sys.stderr)
                return 1
            units.setdefault(code, {"code": code, "name": name, "points": points})
            if code not in codes:
                codes.append(code)
        lists[faction] = codes
    items, item_lists = read_items()
    data = {
        "source": "content/units.xlsx, content/items.xlsx",
        "factions": FACTIONS,
        "units": sorted(units.values(), key=lambda u: u["name"]),
        "lists": lists,
        "items": items,
        # NEUTRAL: every faction; other keys: only warbands of that faction.
        "itemLists": item_lists,
    }
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(units)} characters, {len(lists)} lists, {len(items)} items -> {OUT.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
